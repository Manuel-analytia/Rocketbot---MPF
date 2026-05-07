"""Prologo auto-contenido inyectado al inicio de cada bot hijo.

Resuelve OneDrive, paths, semana forzada (pruebas.txt), credenciales SMTP.
Ver _dev/conventions.md para detalles.

Para inyectar en un bot via parche:
    from pathlib import Path
    PROLOGUE = (Path(__file__).parent / 'autodetect_prologue.py').read_text('utf-8')
    # ... pegar PROLOGUE al inicio del nuevo script
"""

# AUTODETECT_PROLOGUE: resolver todas las rutas desde expanduser(~)
# asi cualquier PC/usuario corre sin necesidad de editar hardvars.
import os as _pre_os
from datetime import datetime as _pre_dt
_home = _pre_os.path.expanduser("~")
_od = None
_marker_sub = "SAN ISIDRO - PROGRAMACION SEMANAL DE EMBARQUES"
try:
    for _nn in _pre_os.listdir(_home):
        if _nn.lower().startswith("onedrive"):
            _cand = _pre_os.path.join(_home, _nn)
            if _pre_os.path.isdir(_pre_os.path.join(_cand, _marker_sub)):
                _od = _cand
                break
except Exception:
    pass
if not _od:
    # Fallback a los paths conocidos
    _od = _pre_os.path.join(_home, "OneDrive - mpf.com.pe (1)")
    if not _pre_os.path.isdir(_od):
        _od = _pre_os.path.join(_home, "OneDrive - mpf.com.pe")
SetVar("v_root_onedrive", _od)
_re = _pre_os.path.join(_od, "SAN ISIDRO - PROGRAMACION SEMANAL DE EMBARQUES")
SetVar("v_root_embarques", _re)
SetVar("v_ruta_base", _re)
SetVar("v_root_riverside", _pre_os.path.join(_od, "SAN ISIDRO - RIVERSIDE NATURAL FOODS, LTD"))
_br = _pre_os.path.join(_re, "0. Bot")
SetVar("v_bot_root", _br)
SetVar("v_bot_logs", _pre_os.path.join(_br, "Logs"))
SetVar("v_bot_temp", _pre_os.path.join(_br, "Temp"))
SetVar("v_bot_rutas", _pre_os.path.join(_br, "Rutas"))
SetVar("v_bot_correos", _pre_os.path.join(_br, "Correos"))
SetVar("v_bot_pruebas", _pre_os.path.join(_br, "Pruebas"))
SetVar("v_ruta_correos_coas", _pre_os.path.join(_br, "Correos", "COAS_e_Informes"))
SetVar("v_ruta_correos_invoice", _pre_os.path.join(_br, "Correos", "Invoice"))
SetVar("v_ruta_correos_estatus", _pre_os.path.join(_br, "Correos", "Estatus"))
SetVar("v_ruta_config_rutas", _pre_os.path.join(_br, "Rutas", "Config_Rutas.xlsx"))
_anio_pre = str(_pre_dt.now().isocalendar()[0])
_semana_pre = str(_pre_dt.now().isocalendar()[1])
# Override con año y semana forzados de pruebas.txt (linea 'N|YYYY')
try:
    _pruebas_pre = _pre_os.path.join(_br, 'Pruebas', 'pruebas.txt')
    if _pre_os.path.exists(_pruebas_pre):
        with open(_pruebas_pre, 'r', encoding='utf-8') as _pfp:
            for _pl in _pfp:
                _pl = _pl.strip()
                if not _pl:
                    continue
                if '|' in _pl:
                    _sw, _sep, _sa = _pl.partition('|')
                    _sw = _sw.strip(); _sa = _sa.strip()
                    if _sw.isdigit() and int(_sw) > 0:
                        _semana_pre = _sw
                        if _sa.isdigit() and int(_sa) >= 2000:
                            _anio_pre = _sa
                        break
                elif _pl.isdigit() and int(_pl) > 0:
                    _semana_pre = _pl
                    break
except Exception:
    pass
# Propagar forzados a vars para que el codigo principal los lea
SetVar("v_anio_semana", _anio_pre)
SetVar("v_numero_semana_trabajo", _semana_pre)
SetVar("v_nombre_semana", "Semana " + _semana_pre)
SetVar("v_ruta_registro_coordinacion", _pre_os.path.join(_re, "COORDINACION DE EMBARQUES",
                                                          "Registro de COAS e informes completos " + _anio_pre + ".xlsx"))

# Leer credenciales SMTP desde Config_Rutas.xlsx (hoja CREDENCIALES)
_config_path = _pre_os.path.join(_br, "Rutas", "Config_Rutas.xlsx")
if _pre_os.path.isfile(_config_path):
    try:
        from openpyxl import load_workbook as _pre_lw
        _wb_cfg = _pre_lw(_config_path, data_only=True)
        if "CREDENCIALES" in _wb_cfg.sheetnames:
            _ws_cred = _wb_cfg["CREDENCIALES"]
            for _r in _ws_cred.iter_rows(min_row=2, values_only=True):
                if not _r or not _r[0]:
                    continue
                _k = str(_r[0]).strip()
                _v = str(_r[1] or "").strip()
                if _k == "SMTP_HOST" and _v: SetVar("v_smtp_host", _v)
                elif _k == "SMTP_PORT" and _v: SetVar("v_smtp_port", _v)
                elif _k == "SMTP_USER" and _v: SetVar("v_smtp_user", _v)
                elif _k == "SMTP_PASS" and _v: SetVar("v_smtp_pass", _v)
                elif _k == "SMTP_SSL" and _v: SetVar("v_smtp_ssl", _v.lower())
    except Exception:
        pass
# Path para libs locales del bot (pypdf, etc. instalados sin admin)
try:
    import sys as _sys_pre
    _libs_pre = _pre_os.path.join(_br, 'libs')
    if _pre_os.path.isdir(_libs_pre) and _libs_pre not in _sys_pre.path:
        _sys_pre.path.insert(0, _libs_pre)
except Exception:
    pass
# FIN AUTODETECT_PROLOGUE
