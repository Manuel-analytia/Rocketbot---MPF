# AUTODETECT_PROLOGUE: resolver todas las rutas desde expanduser(~)
# asi cualquier PC/usuario corre sin necesidad de editar hardvars.
import os as _pre_os
from datetime import datetime as _pre_dt
_home = _pre_os.path.expanduser("~")
_od = None
_marker_sub = "SAN ISIDRO - PROGRAMACION SEMANAL DE EMBARQUES"
try:
    for _nn in _pre_os.listdir(_home):
        if _nn.lower().startswith("onedrive"):
            _cand = _pre_os.path.join(_home, _nn)
            if _pre_os.path.isdir(_pre_os.path.join(_cand, _marker_sub)):
                _od = _cand
                break
except Exception:
    pass
if not _od:
    # Fallback a los paths conocidos
    _od = _pre_os.path.join(_home, "OneDrive - mpf.com.pe (1)")
    if not _pre_os.path.isdir(_od):
        _od = _pre_os.path.join(_home, "OneDrive - mpf.com.pe")
SetVar("v_root_onedrive", _od)
_re = _pre_os.path.join(_od, "SAN ISIDRO - PROGRAMACION SEMANAL DE EMBARQUES")
SetVar("v_root_embarques", _re)
SetVar("v_ruta_base", _re)
SetVar("v_root_riverside", _pre_os.path.join(_od, "SAN ISIDRO - RIVERSIDE NATURAL FOODS, LTD"))
_br = _pre_os.path.join(_re, "0. Bot")
SetVar("v_bot_root", _br)
SetVar("v_bot_logs", _pre_os.path.join(_br, "Logs"))
SetVar("v_bot_temp", _pre_os.path.join(_br, "Temp"))
SetVar("v_bot_rutas", _pre_os.path.join(_br, "Rutas"))
SetVar("v_bot_correos", _pre_os.path.join(_br, "Correos"))
SetVar("v_bot_pruebas", _pre_os.path.join(_br, "Pruebas"))
SetVar("v_ruta_correos_coas", _pre_os.path.join(_br, "Correos", "COAS_e_Informes"))
SetVar("v_ruta_correos_invoice", _pre_os.path.join(_br, "Correos", "Invoice"))
SetVar("v_ruta_correos_estatus", _pre_os.path.join(_br, "Correos", "Estatus"))
SetVar("v_ruta_config_rutas", _pre_os.path.join(_br, "Rutas", "Config_Rutas.xlsx"))
_anio_pre = str(_pre_dt.now().isocalendar()[0])
_semana_pre = str(_pre_dt.now().isocalendar()[1])
# Override con año y semana forzados de pruebas.txt (linea 'N|YYYY')
try:
    _pruebas_pre = _pre_os.path.join(_br, 'Pruebas', 'pruebas.txt')
    if _pre_os.path.exists(_pruebas_pre):
        with open(_pruebas_pre, 'r', encoding='utf-8') as _pfp:
            for _pl in _pfp:
                _pl = _pl.strip()
                if not _pl:
                    continue
                if '|' in _pl:
                    _sw, _sep, _sa = _pl.partition('|')
                    _sw = _sw.strip(); _sa = _sa.strip()
                    if _sw.isdigit() and int(_sw) > 0:
                        _semana_pre = _sw
                        if _sa.isdigit() and int(_sa) >= 2000:
                            _anio_pre = _sa
                        break
                elif _pl.isdigit() and int(_pl) > 0:
                    _semana_pre = _pl
                    break
except Exception:
    pass
# Propagar forzados a vars para que el codigo principal los lea
SetVar("v_anio_semana", _anio_pre)
SetVar("v_numero_semana_trabajo", _semana_pre)
SetVar("v_nombre_semana", "Semana " + _semana_pre)
SetVar("v_ruta_registro_coordinacion", _pre_os.path.join(_re, "COORDINACION DE EMBARQUES",
                                                          "Registro de COAS e informes completos " + _anio_pre + ".xlsx"))

# Leer credenciales SMTP desde Config_Rutas.xlsx (hoja CREDENCIALES)
_config_path = _pre_os.path.join(_br, "Rutas", "Config_Rutas.xlsx")
if _pre_os.path.isfile(_config_path):
    try:
        from openpyxl import load_workbook as _pre_lw
        _wb_cfg = _pre_lw(_config_path, data_only=True)
        if "CREDENCIALES" in _wb_cfg.sheetnames:
            _ws_cred = _wb_cfg["CREDENCIALES"]
            for _r in _ws_cred.iter_rows(min_row=2, values_only=True):
                if not _r or not _r[0]:
                    continue
                _k = str(_r[0]).strip()
                _v = str(_r[1] or "").strip()
                if _k == "SMTP_HOST" and _v: SetVar("v_smtp_host", _v)
                elif _k == "SMTP_PORT" and _v: SetVar("v_smtp_port", _v)
                elif _k == "SMTP_USER" and _v: SetVar("v_smtp_user", _v)
                elif _k == "SMTP_PASS" and _v: SetVar("v_smtp_pass", _v)
                elif _k == "SMTP_SSL" and _v: SetVar("v_smtp_ssl", _v.lower())
    except Exception:
        pass
# Path para libs locales del bot (pypdf, etc. instalados sin admin)
try:
    import sys as _sys_pre
    _libs_pre = _pre_os.path.join(_br, 'libs')
    if _pre_os.path.isdir(_libs_pre) and _libs_pre not in _sys_pre.path:
        _sys_pre.path.insert(0, _libs_pre)
except Exception:
    pass
# FIN AUTODETECT_PROLOGUE

from datetime import datetime
import os
import re
import zipfile
from xml.sax.saxutils import escape
from xml.etree import ElementTree as ET

PROCESO = '0_1_Verificacion_Condiciones'
NS_MAIN = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
NS_REL = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
NS_PKG = 'http://schemas.openxmlformats.org/package/2006/relationships'


def hora_txt():
    return datetime.now().strftime('%H:%M:%S')


def crear_carpeta_segura(ruta):
    os.makedirs(ruta, exist_ok=True)
    return ruta


def escribir_linea_log(evento):
    ruta_log = GetVar('v_ruta_log_actual')
    if not ruta_log:
        return
    try:
        crear_carpeta_segura(os.path.dirname(ruta_log))
        with open(ruta_log, 'a', encoding='utf-8') as archivo:
            archivo.write(f"{hora_txt()}|{PROCESO}|{evento}\n")
    except Exception:
        pass


def comentario_archivo(ruta, contenido):
    if not os.path.exists(ruta):
        with open(ruta, 'w', encoding='utf-8') as archivo:
            archivo.write(contenido)


def col_excel(numero):
    resultado = ''
    while numero > 0:
        numero, resto = divmod(numero - 1, 26)
        resultado = chr(65 + resto) + resultado
    return resultado


def crear_xlsx_stdlib(ruta, hojas):
    from xml.sax.saxutils import escape
    import zipfile
    def col_excel(numero):
        resultado = ""
        while numero > 0:
            numero, resto = divmod(numero - 1, 26)
            resultado = chr(65 + resto) + resultado
        return resultado

    def xml_hoja(filas):
        partes = [
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>',
            '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">',
            '<sheetData>'
        ]
        for idx_fila, fila in enumerate(filas, start=1):
            partes.append(f'<row r="{idx_fila}">')
            for idx_col, valor in enumerate(fila, start=1):
                ref = f'{col_excel(idx_col)}{idx_fila}'
                texto = '' if valor is None else str(valor)
                partes.append(f'<c r="{ref}" t="inlineStr"><is><t>{escape(texto)}</t></is></c>')
            partes.append('</row>')
        partes.extend(['</sheetData>', '</worksheet>'])
        return ''.join(partes)

    nombres = list(hojas.keys())
    contenido_types = [
        '<?xml version="1.0" encoding="UTF-8"?>',
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">',
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>',
        '<Default Extension="xml" ContentType="application/xml"/>',
        '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>',
        '<Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>'
    ]
    for i in range(1, len(nombres) + 1):
        contenido_types.append(f'<Override PartName="/xl/worksheets/sheet{i}.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>')
    contenido_types.append('</Types>')

    rels_root = '<?xml version="1.0" encoding="UTF-8"?>\n<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">\n  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>\n</Relationships>'

    workbook = [
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>',
        '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">',
        '<sheets>'
    ]
    for i, nombre in enumerate(nombres, start=1):
        workbook.append(f'<sheet name="{escape(nombre)}" sheetId="{i}" r:id="rId{i}"/>')
    workbook.extend(['</sheets>', '</workbook>'])

    workbook_rels = ['<?xml version="1.0" encoding="UTF-8"?>', '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">']
    for i in range(1, len(nombres) + 1):
        workbook_rels.append(f'<Relationship Id="rId{i}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet{i}.xml"/>')
    workbook_rels.append(f'<Relationship Id="rId{len(nombres) + 1}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>')
    workbook_rels.append('</Relationships>')

    styles = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">\n  <fonts count="1"><font><sz val="11"/><name val="Calibri"/></font></fonts>\n  <fills count="1"><fill><patternFill patternType="none"/></fill></fills>\n  <borders count="1"><border/></borders>\n  <cellStyleXfs count="1"><xf/></cellStyleXfs>\n  <cellXfs count="1"><xf xfId="0"/></cellXfs>\n  <cellStyles count="1"><cellStyle name="Normal" xfId="0" builtinId="0"/></cellStyles>\n</styleSheet>'

    with zipfile.ZipFile(ruta, 'w', compression=zipfile.ZIP_DEFLATED) as zf:
        zf.writestr('[Content_Types].xml', ''.join(contenido_types))
        zf.writestr('_rels/.rels', rels_root)
        zf.writestr('xl/workbook.xml', ''.join(workbook))
        zf.writestr('xl/_rels/workbook.xml.rels', ''.join(workbook_rels))
        zf.writestr('xl/styles.xml', styles)
        for i, nombre in enumerate(nombres, start=1):
            zf.writestr(f'xl/worksheets/sheet{i}.xml', xml_hoja(hojas[nombre]))


def crear_config_rutas_xlsx(ruta):
    hojas = {
        'BASE': [
            ['PARAMETRO', 'VALOR', 'OBLIGATORIO', 'DESCRIPCION'],
            ['ROOT_EMBARQUES', 'SAN ISIDRO - PROGRAMACION SEMANAL DE EMBARQUES', 'SI', 'Raíz de embarques (bajo ROOT_ONEDRIVE)'],
            ['ROOT_RIVERSIDE', 'SAN ISIDRO - RIVERSIDE NATURAL FOODS, LTD', 'SI', 'Raíz del cliente Riverside (bajo ROOT_ONEDRIVE)'],
            ['ROOT_COAS_FUENTE', 'Archivos de Juliana Legua - ORGANICO', 'SI', 'Carpeta fuente de los PDF de COAs (bajo ROOT_ONEDRIVE)'],
            ['ROOT_INFORMES_PREFIJO', 'PISCO - ', 'SI', 'Prefijo de la carpeta de informes; el bot concatena el año ISO']
        ],
        'SUBRUTAS_BOT': [
            ['PARAMETRO', 'VALOR', 'OBLIGATORIO', 'DESCRIPCION'],
            ['BOT_ROOT', '0. Bot', 'SI', 'Carpeta técnica central, fija por diseño'],
            ['BOT_LOGS', 'Logs', 'SI', 'Subcarpeta de bitácoras'],
            ['BOT_RUTAS', 'Rutas', 'SI', 'Subcarpeta de configuración'],
            ['BOT_CORREOS', 'Correos', 'SI', 'Subcarpeta de destinatarios TXT'],
            ['BOT_PRUEBAS', 'Pruebas', 'SI', 'Subcarpeta de pruebas controladas'],
            ['BOT_TEMP', 'Temp', 'SI', 'Subcarpeta temporal'],
            ['BOT_EVIDENCIAS', 'Evidencias', 'NO', 'Subcarpeta opcional de trazabilidad']
        ],
        'SUBRUTAS_NEGOCIO': [
            ['PARAMETRO', 'VALOR', 'OBLIGATORIO', 'DESCRIPCION'],
            ['CARPETA_COORDINACION_EMBARQUES', 'COORDINACION DE EMBARQUES', 'SI', 'Subcarpeta donde vive el Registro Coordinacion'],
            ['NOMBRE_REGISTRO_PREFIJO', 'Registro de COAS e informes completos ', 'SI', 'Prefijo del archivo del Registro; el bot agrega {AÑO}.xlsx'],
            ['PAGINA_EXCLUIR_INFORMES', 'TERMINOS Y CONDICIONES', 'SI', 'Si una página del PDF de informes contiene este texto, se omite al consolidar']
        ],
        'CREDENCIALES': [
            ['PARAMETRO', 'VALOR', 'OBLIGATORIO', 'DESCRIPCION'],
            ['SMTP_HOST', 'smtp.office365.com', 'SI', 'Servidor SMTP'],
            ['SMTP_PORT', '587', 'SI', 'Puerto SMTP'],
            ['SMTP_USER', '', 'SI', 'Usuario remitente'],
            ['SMTP_PASS', '', 'SI', 'Clave SMTP'],
            ['SMTP_SSL', 'false', 'NO', 'Usar SSL en SMTP (false/true)'],
        ]
    }
    try:
        from openpyxl import Workbook
        wb = Workbook()
        primera = True
        for nombre, filas in hojas.items():
            if primera:
                ws = wb.active
                ws.title = nombre
                primera = False
            else:
                ws = wb.create_sheet(title=nombre)
            for fila in filas:
                ws.append(fila)
        wb.save(ruta)
    except Exception:
        crear_xlsx_stdlib(ruta, hojas)


def leer_xlsx_stdlib(ruta):
    import zipfile
    from xml.etree import ElementTree as ET
    _ns_main = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
    _ns_rel  = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
    _ns_pkg  = 'http://schemas.openxmlformats.org/package/2006/relationships'
    ns = {'m': _ns_main, 'r': _ns_rel, 'p': _ns_pkg}

    def texto_celda(celda, shared):
        tipo = celda.attrib.get('t', '')
        if tipo == 'inlineStr':
            return ''.join([(nodo.text or '') for nodo in celda.findall('.//m:t', ns)])
        valor = celda.find('m:v', ns)
        if valor is None:
            return ''
        if tipo == 's':
            try:
                idx = int(valor.text)
                return shared[idx] if idx < len(shared) else ''
            except Exception:
                return valor.text or ''
        return valor.text or ''

    hojas = {}
    with zipfile.ZipFile(ruta, 'r') as zf:
        shared = []
        if 'xl/sharedStrings.xml' in zf.namelist():
            raiz_shared = ET.fromstring(zf.read('xl/sharedStrings.xml'))
            for si in raiz_shared.findall('m:si', ns):
                shared.append(''.join([(nodo.text or '') for nodo in si.findall('.//m:t', ns)]))

        workbook = ET.fromstring(zf.read('xl/workbook.xml'))
        rels = ET.fromstring(zf.read('xl/_rels/workbook.xml.rels'))
        rel_map = {}
        for rel in rels.findall('p:Relationship', ns):
            rel_map[rel.attrib.get('Id')] = rel.attrib.get('Target')

        for sheet in workbook.findall('m:sheets/m:sheet', ns):
            nombre = sheet.attrib.get('name', '')
            rid = sheet.attrib.get('{%s}id' % NS_REL)
            target = rel_map.get(rid, '')
            if not target:
                continue
            path = 'xl/' + target.lstrip('/')
            raiz_hoja = ET.fromstring(zf.read(path))
            filas = []
            for row in raiz_hoja.findall('m:sheetData/m:row', ns):
                valores = []
                for celda in row.findall('m:c', ns):
                    valores.append(texto_celda(celda, shared))
                filas.append(valores)
            hojas[nombre] = filas
    return hojas


def leer_config_rutas_xlsx(ruta):
    hojas = None
    try:
        from openpyxl import load_workbook
        wb = load_workbook(ruta, data_only=True)
        hojas = {}
        for nombre in wb.sheetnames:
            ws = wb[nombre]
            filas = []
            for fila in ws.iter_rows(values_only=True):
                filas.append(['' if valor is None else str(valor) for valor in fila])
            hojas[nombre] = filas
    except Exception:
        hojas = leer_xlsx_stdlib(ruta)

    resultado = {}
    for nombre_hoja, filas in hojas.items():
        data_hoja = {}
        for fila in filas[1:]:
            clave = fila[0].strip() if len(fila) > 0 and fila[0] else ''
            valor = fila[1].strip() if len(fila) > 1 and fila[1] else ''
            if clave:
                data_hoja[clave] = valor
        resultado[nombre_hoja] = data_hoja
    return resultado


def resolver_ruta(base, valor):
    valor = (valor or '').strip()
    if not valor:
        return ''
    if re.match(r'^[A-Za-z]:\\', valor) or valor.startswith('\\'):
        return os.path.normpath(valor)
    return os.path.normpath(os.path.join(base, valor))


SetVar('v_precheck_estado', 'EN_PROCESO')
SetVar('v_precheck_detalle', 'Precheck iniciado')

errores = []
advertencias = []

root_onedrive_default = GetVar('v_root_onedrive')
root_embarques_default = 'SAN ISIDRO - PROGRAMACION SEMANAL DE EMBARQUES'
root_riverside_default = 'SAN ISIDRO - RIVERSIDE NATURAL FOODS, LTD'

root_embarques_inicial = resolver_ruta(root_onedrive_default, root_embarques_default)
root_riverside_inicial = resolver_ruta(root_onedrive_default, root_riverside_default)
bot_root_fijo = os.path.join(root_embarques_inicial, '0. Bot')
bot_logs_fija = os.path.join(bot_root_fijo, 'Logs')
bot_rutas_fija = os.path.join(bot_root_fijo, 'Rutas')

SetVar('v_root_onedrive', root_onedrive_default)
SetVar('v_root_embarques', root_embarques_inicial)
SetVar('v_root_riverside', root_riverside_inicial)
SetVar('v_bot_root', bot_root_fijo)
SetVar('v_bot_logs', bot_logs_fija)
SetVar('v_bot_rutas', bot_rutas_fija)
SetVar('v_ruta_config_rutas', os.path.join(bot_rutas_fija, 'Config_Rutas.xlsx'))
SetVar('v_ruta_log_actual', os.path.join(bot_logs_fija, datetime.now().strftime('LOG - %Y-%m-%d %H.%M.%S.txt')))
SetVar('v_orq_log_path', GetVar('v_ruta_log_actual'))

if not os.path.exists(root_onedrive_default):
    errores.append(f'No existe ROOT_ONEDRIVE: {root_onedrive_default}')

if not errores:
    if not os.path.exists(root_embarques_inicial):
        errores.append(f'No existe ROOT_EMBARQUES: {root_embarques_inicial}')
    if not os.path.exists(root_riverside_inicial):
        errores.append(f'No existe ROOT_RIVERSIDE: {root_riverside_inicial}')

if not errores:
    for carpeta in [bot_root_fijo, bot_logs_fija, bot_rutas_fija, os.path.join(bot_root_fijo, 'Correos'), os.path.join(bot_root_fijo, 'Pruebas'), os.path.join(bot_root_fijo, 'Temp')]:
        crear_carpeta_segura(carpeta)
    escribir_linea_log('Inicio del proceso')
    escribir_linea_log('Estructura base 0. Bot validada o creada')

config_path = GetVar('v_ruta_config_rutas')
if not errores:
    if not os.path.exists(config_path):
        crear_config_rutas_xlsx(config_path)
        escribir_linea_log('Config_Rutas.xlsx creado con plantilla base')

    ruta_correos = os.path.join(bot_root_fijo, 'Correos')
    ruta_pruebas = os.path.join(bot_root_fijo, 'Pruebas')
    for _sub in ('Invoice', 'COAS_e_Informes', 'Estatus'):
        _subdir = os.path.join(ruta_correos, _sub)
        os.makedirs(_subdir, exist_ok=True)
        comentario_archivo(os.path.join(_subdir, 'Para.txt'), '# Un correo por linea\n')
        comentario_archivo(os.path.join(_subdir, 'CC.txt'), '# Un correo por linea\n')
        comentario_archivo(os.path.join(_subdir, 'CO.txt'), '# Un correo por linea\n')
    _dt_pre = __import__('datetime').datetime.now().isocalendar()
    comentario_archivo(os.path.join(ruta_pruebas, 'pruebas.txt'), 'semana|anio\n' + str(_dt_pre[1]) + '|' + str(_dt_pre[0]) + '\n')
    escribir_linea_log('Archivos base de Correos (Invoice/COAS_e_Informes/Estatus) y Pruebas validados')

config = {}
if not errores and os.path.exists(config_path):
    try:
        config = leer_config_rutas_xlsx(config_path)
    except Exception as exc:
        errores.append(f'No se pudo leer Config_Rutas.xlsx: {exc}')

if not errores:
    requeridos = [
        ('BASE', 'ROOT_EMBARQUES'),
        ('BASE', 'ROOT_RIVERSIDE'),
        ('SUBRUTAS_BOT', 'BOT_ROOT'),
        ('SUBRUTAS_BOT', 'BOT_LOGS'),
        ('SUBRUTAS_BOT', 'BOT_RUTAS'),
        ('SUBRUTAS_BOT', 'BOT_CORREOS'),
        ('SUBRUTAS_BOT', 'BOT_PRUEBAS'),
        ('SUBRUTAS_BOT', 'BOT_TEMP'),
        ('BASE', 'ROOT_COAS_FUENTE'),
        ('BASE', 'ROOT_INFORMES_PREFIJO'),
        ('SUBRUTAS_NEGOCIO', 'CARPETA_COORDINACION_EMBARQUES'),
        ('SUBRUTAS_NEGOCIO', 'NOMBRE_REGISTRO_PREFIJO'),
        ('SUBRUTAS_NEGOCIO', 'PAGINA_EXCLUIR_INFORMES'),
        ('CREDENCIALES', 'SMTP_HOST'),
        ('CREDENCIALES', 'SMTP_PORT'),
        ('CREDENCIALES', 'SMTP_USER'),
        ('CREDENCIALES', 'SMTP_PASS'),
    ]
    for hoja, clave in requeridos:
        valor = config.get(hoja, {}).get(clave, '').strip()
        if not valor:
            errores.append(f'Campo obligatorio vacío en Config_Rutas.xlsx: {hoja}.{clave}')

if not errores:
    base_cfg = config.get('BASE', {})
    bot_cfg = config.get('SUBRUTAS_BOT', {})
    neg_cfg = config.get('SUBRUTAS_NEGOCIO', {})

    root_onedrive = root_onedrive_default  # siempre desde v_root_onedrive resuelto por el padre
    root_embarques_val = base_cfg.get('ROOT_EMBARQUES', root_embarques_default).strip()
    root_riverside_val = base_cfg.get('ROOT_RIVERSIDE', root_riverside_default).strip()

    bot_root_cfg = (bot_cfg.get('BOT_ROOT', '0. Bot') or '0. Bot').strip()
    if bot_root_cfg != '0. Bot':
        advertencias.append(f'BOT_ROOT fue normalizado a 0. Bot (valor recibido: {bot_root_cfg})')
    bot_root_name = '0. Bot'
    bot_logs_name = bot_cfg.get('BOT_LOGS', 'Logs').strip()
    bot_rutas_name = bot_cfg.get('BOT_RUTAS', 'Rutas').strip()
    bot_correos_name = bot_cfg.get('BOT_CORREOS', 'Correos').strip()
    bot_pruebas_name = bot_cfg.get('BOT_PRUEBAS', 'Pruebas').strip()
    bot_temp_name = bot_cfg.get('BOT_TEMP', 'Temp').strip()
    bot_evidencias_name = bot_cfg.get('BOT_EVIDENCIAS', '').strip()

    root_embarques = resolver_ruta(root_onedrive, root_embarques_val)
    root_riverside = resolver_ruta(root_onedrive, root_riverside_val)

    if not os.path.exists(root_onedrive):
        errores.append(f'No existe ROOT_ONEDRIVE configurado: {root_onedrive}')
    if not os.path.exists(root_embarques):
        errores.append(f'No existe ROOT_EMBARQUES configurado: {root_embarques}')
    if not os.path.exists(root_riverside):
        errores.append(f'No existe ROOT_RIVERSIDE configurado: {root_riverside}')

    bot_root = os.path.join(root_embarques, bot_root_name)
    bot_logs = os.path.join(bot_root, bot_logs_name)
    bot_rutas = os.path.join(bot_root, bot_rutas_name)
    bot_correos = os.path.join(bot_root, bot_correos_name)
    bot_pruebas = os.path.join(bot_root, bot_pruebas_name)
    bot_temp = os.path.join(bot_root, bot_temp_name)
    bot_evidencias = os.path.join(bot_root, bot_evidencias_name) if bot_evidencias_name else ''
    ruta_config = os.path.join(bot_rutas, 'Config_Rutas.xlsx')
    ruta_para = os.path.join(bot_correos, 'Para.txt')
    ruta_cc = os.path.join(bot_correos, 'CC.txt')
    ruta_co = os.path.join(bot_correos, 'CO.txt')
    ruta_pruebas_txt = os.path.join(bot_pruebas, 'pruebas.txt')

    SetVar('v_root_onedrive', root_onedrive)
    SetVar('v_root_embarques', root_embarques)
    SetVar('v_root_riverside', root_riverside)
    SetVar('v_bot_root', bot_root)
    SetVar('v_bot_logs', bot_logs)
    SetVar('v_bot_rutas', bot_rutas)
    SetVar('v_bot_correos', bot_correos)
    SetVar('v_bot_pruebas', bot_pruebas)
    SetVar('v_bot_temp', bot_temp)
    SetVar('v_bot_evidencias', bot_evidencias)
    SetVar('v_ruta_config_rutas', ruta_config)
    SetVar('v_ruta_para_txt', ruta_para)
    SetVar('v_ruta_cc_txt', ruta_cc)
    SetVar('v_ruta_co_txt', ruta_co)
    SetVar('v_ruta_pruebas_txt', ruta_pruebas_txt)
    SetVar('v_carpeta_coordinacion_embarques', neg_cfg.get('CARPETA_COORDINACION_EMBARQUES', '').strip())
    SetVar('v_nombre_registro_prefijo', neg_cfg.get('NOMBRE_REGISTRO_PREFIJO', '').strip())
    SetVar('v_pagina_excluir_informes', neg_cfg.get('PAGINA_EXCLUIR_INFORMES', '').strip())
    SetVar('v_root_coas_fuente', base_cfg.get('ROOT_COAS_FUENTE', '').strip())
    SetVar('v_root_informes_prefijo', base_cfg.get('ROOT_INFORMES_PREFIJO', '').strip())

    cred_cfg = config.get('CREDENCIALES', {})
    SetVar('v_smtp_host', cred_cfg.get('SMTP_HOST', '').strip())
    SetVar('v_smtp_port', cred_cfg.get('SMTP_PORT', '').strip())
    SetVar('v_smtp_user', cred_cfg.get('SMTP_USER', '').strip())
    SetVar('v_smtp_pass', cred_cfg.get('SMTP_PASS', '').strip())
    SetVar('v_smtp_ssl', (cred_cfg.get('SMTP_SSL', 'false') or 'false').strip().lower())

    if not errores:
        for carpeta in [bot_root, bot_logs, bot_rutas, bot_correos, bot_pruebas, bot_temp]:
            crear_carpeta_segura(carpeta)
        if bot_evidencias:
            crear_carpeta_segura(bot_evidencias)
        SetVar('v_ruta_log_actual', os.path.join(bot_logs, datetime.now().strftime('LOG - %Y-%m-%d %H.%M.%S.txt')))
        SetVar('v_orq_log_path', GetVar('v_ruta_log_actual'))
        escribir_linea_log('Ruta ROOT_EMBARQUES válida')
        escribir_linea_log('Ruta ROOT_RIVERSIDE válida')
        if advertencias:
            for advertencia in advertencias:
                escribir_linea_log(advertencia)

        libs_bot = os.path.join(bot_root, 'libs')
        SetVar('v_bot_libs', libs_bot)

        # Los Para/CC/CO ahora se crean solo en subcarpetas (Invoice/COAS/Estatus)
        _dt_init = __import__('datetime').datetime.now().isocalendar()
        comentario_archivo(ruta_pruebas_txt, 'semana|anio\n' + str(_dt_init[1]) + '|' + str(_dt_init[0]) + '\n')

        # Subcarpetas separadas por tipo de correo (post reunion Karla 2026-04-22)
        ruta_correos_coas = os.path.join(bot_correos, 'COAS_e_Informes')
        ruta_correos_invoice = os.path.join(bot_correos, 'Invoice')
        ruta_correos_estatus = os.path.join(bot_correos, 'Estatus')
        crear_carpeta_segura(ruta_correos_coas)
        crear_carpeta_segura(ruta_correos_invoice)
        crear_carpeta_segura(ruta_correos_estatus)
        SetVar('v_ruta_correos_coas', ruta_correos_coas)
        SetVar('v_ruta_correos_invoice', ruta_correos_invoice)
        SetVar('v_ruta_correos_estatus', ruta_correos_estatus)
        comentario_archivo(os.path.join(ruta_correos_coas, 'Para.txt'), '# Un correo por línea\n')
        comentario_archivo(os.path.join(ruta_correos_coas, 'CC.txt'), '# Un correo por línea\n')
        comentario_archivo(os.path.join(ruta_correos_coas, 'CO.txt'), '# Un correo por línea\n')
        comentario_archivo(os.path.join(ruta_correos_invoice, 'Para.txt'), '# Un correo por línea\nmcinvoicecapture@concursolutions.com\n')
        comentario_archivo(os.path.join(ruta_correos_invoice, 'CC.txt'), '# Un correo por línea\n')
        comentario_archivo(os.path.join(ruta_correos_invoice, 'CO.txt'), '# Un correo por línea\n')
        escribir_linea_log('Subcarpetas Correos/COAS_e_Informes y Correos/Invoice validadas')

        # Registro Coordinacion del año actual (post reunion Karla 2026-04-22)
        registro_dir = os.path.join(root_embarques, neg_cfg.get('CARPETA_COORDINACION_EMBARQUES', '').strip())
        registro_prefijo = neg_cfg.get('NOMBRE_REGISTRO_PREFIJO', '').strip()
        _anio_reg = datetime.now().isocalendar()[0]
        # Asegurar espacio entre prefijo y anio (el lector de xlsx strippea trailing space)
        _prefijo_norm = registro_prefijo.rstrip() + ' ' if registro_prefijo else ''
        registro_nombre = _prefijo_norm + str(_anio_reg) + '.xlsx'
        registro_path = os.path.join(registro_dir, registro_nombre)
        SetVar('v_ruta_registro_coordinacion', registro_path)
        if registro_dir and not os.path.exists(registro_path):
            try:
                crear_carpeta_segura(registro_dir)
                header = ['Cliente', 'Embarque', 'Fecha', 'N° de pedido', 'COAS', 'Informes',
                         'Estado Invoice', 'Fecha de envío Invoice', 'Existencia', 'Fecha de envío COAs']
                meses = ['ENE', 'FEB', 'MAR', 'ABR', 'MAY', 'JUN',
                         'JUL', 'AGO', 'SEP', 'OCT', 'NOV', 'DIC']
                hojas_reg = {m: [header] for m in meses}
                crear_xlsx_stdlib(registro_path, hojas_reg)
                escribir_linea_log('Registro Coordinacion creado: ' + registro_path)
            except Exception as exc_reg:
                advertencias.append('No se pudo crear Registro Coordinacion: ' + str(exc_reg))

        ruta_envios = os.path.join(bot_logs, 'Envios.xlsx')
        SetVar('v_ruta_envios', ruta_envios)
        if not os.path.exists(ruta_envios):
            try:
                crear_xlsx_stdlib(ruta_envios, {
                    'ENVIOS': [['FECHA', 'HORA', 'EMBARQUE', 'PEDIDO', 'CLIENTE', 'TIPO_DOC', 'PO', 'ARCHIVO', 'ESTADO', 'DETALLE']]
                })
                escribir_linea_log('Envios.xlsx creado con encabezados')
            except Exception as exc_env:
                errores.append(f'No se pudo crear Envios.xlsx: {exc_env}')

        semana_forzada = ''
        semana_actual = datetime.now().isocalendar()[1]
        anio_actual = datetime.now().isocalendar()[0]
        try:
            with open(ruta_pruebas_txt, 'r', encoding='utf-8') as archivo:
                for _pl in archivo:
                    _pl = _pl.strip()
                    if not _pl:
                        continue
                    if '|' in _pl:
                        _sw, _sep, _sa = _pl.partition('|')
                        _sw = _sw.strip(); _sa = _sa.strip()
                        if _sw.isdigit() and int(_sw) > 0:
                            semana_forzada = _sw
                            if _sa.isdigit() and int(_sa) >= 2000:
                                anio_forzado = _sa
                    elif _pl.isdigit() and int(_pl) > 0:
                        semana_forzada = _pl
        except Exception:
            semana_forzada = ''

        numero_semana_trabajo = semana_forzada if semana_forzada else str(semana_actual)
        SetVar('v_semana_forzada', semana_forzada)
        SetVar('v_numero_semana_trabajo', numero_semana_trabajo)
        SetVar('v_anio_semana', anio_forzado if anio_forzado else str(anio_actual))

if errores:
    detalle = 'ERROR | ' + ' ; '.join(errores)
    SetVar('v_precheck_estado', 'ERROR')
    SetVar('v_precheck_detalle', detalle)
    escribir_linea_log(detalle)
else:
    resumen = [
        'OK',
        f"ROOT_EMBARQUES={GetVar('v_root_embarques')}",
        f"ROOT_RIVERSIDE={GetVar('v_root_riverside')}",
        f"BOT_ROOT={GetVar('v_bot_root')}",
        f"ANIO={anio_forzado if anio_forzado else anio_actual}",
        f"SEMANA={GetVar('v_numero_semana_trabajo')}",
        f"NOMBRE_SEMANA=Semana {GetVar('v_numero_semana_trabajo')}",
        'RUTA_SEMANA_ESPERADA=' + os.path.join(GetVar('v_root_embarques'), str(anio_forzado if anio_forzado else anio_actual), 'Semana ' + str(GetVar('v_numero_semana_trabajo')))
    ]
    if advertencias:
        resumen.append('ADVERTENCIAS=' + ' | '.join(advertencias))
    detalle = ' | '.join(resumen)
    SetVar('v_precheck_estado', 'OK')
    SetVar('v_precheck_detalle', detalle)
    escribir_linea_log('Precheck finalizado correctamente')
