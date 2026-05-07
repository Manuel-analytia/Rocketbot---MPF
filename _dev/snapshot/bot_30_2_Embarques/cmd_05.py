# AUTODETECT_PROLOGUE
import os as _pre_os
from datetime import datetime as _pre_dt
_home = _pre_os.path.expanduser("~")
_od = None
_marker_sub = "SAN ISIDRO - PROGRAMACION SEMANAL DE EMBARQUES"
try:
    for _nn in _pre_os.listdir(_home):
        if _nn.lower().startswith("onedrive"):
            _cand = _pre_os.path.join(_home, _nn)
            if _pre_os.path.isdir(_pre_os.path.join(_cand, _marker_sub)):
                _od = _cand
                break
except Exception:
    pass
if not _od:
    # Fallback a los paths conocidos
    _od = _pre_os.path.join(_home, "OneDrive - mpf.com.pe (1)")
    if not _pre_os.path.isdir(_od):
        _od = _pre_os.path.join(_home, "OneDrive - mpf.com.pe")
SetVar("v_root_onedrive", _od)
_re = _pre_os.path.join(_od, "SAN ISIDRO - PROGRAMACION SEMANAL DE EMBARQUES")
SetVar("v_root_embarques", _re)
SetVar("v_ruta_base", _re)
SetVar("v_root_riverside", _pre_os.path.join(_od, "SAN ISIDRO - RIVERSIDE NATURAL FOODS, LTD"))
_br = _pre_os.path.join(_re, "0. Bot")
SetVar("v_bot_root", _br)
SetVar("v_bot_logs", _pre_os.path.join(_br, "Logs"))
SetVar("v_bot_temp", _pre_os.path.join(_br, "Temp"))
SetVar("v_bot_rutas", _pre_os.path.join(_br, "Rutas"))
SetVar("v_bot_correos", _pre_os.path.join(_br, "Correos"))
SetVar("v_bot_pruebas", _pre_os.path.join(_br, "Pruebas"))
SetVar("v_ruta_correos_coas", _pre_os.path.join(_br, "Correos", "COAS_e_Informes"))
SetVar("v_ruta_correos_invoice", _pre_os.path.join(_br, "Correos", "Invoice"))
SetVar("v_ruta_correos_estatus", _pre_os.path.join(_br, "Correos", "Estatus"))
SetVar("v_ruta_config_rutas", _pre_os.path.join(_br, "Rutas", "Config_Rutas.xlsx"))
_anio_pre = str(_pre_dt.now().isocalendar()[0])
_semana_pre = str(_pre_dt.now().isocalendar()[1])
# Override con año y semana forzados de pruebas.txt (linea 'N|YYYY')
try:
    _pruebas_pre = _pre_os.path.join(_br, 'Pruebas', 'pruebas.txt')
    if _pre_os.path.exists(_pruebas_pre):
        with open(_pruebas_pre, 'r', encoding='utf-8') as _pfp:
            for _pl in _pfp:
                _pl = _pl.strip()
                if not _pl:
                    continue
                if '|' in _pl:
                    _sw, _sep, _sa = _pl.partition('|')
                    _sw = _sw.strip(); _sa = _sa.strip()
                    if _sw.isdigit() and int(_sw) > 0:
                        _semana_pre = _sw
                        if _sa.isdigit() and int(_sa) >= 2000:
                            _anio_pre = _sa
                        break
                elif _pl.isdigit() and int(_pl) > 0:
                    _semana_pre = _pl
                    break
except Exception:
    pass
# Propagar forzados a vars para que el codigo principal los lea
SetVar("v_anio_semana", _anio_pre)
SetVar("v_numero_semana_trabajo", _semana_pre)
SetVar("v_nombre_semana", "Semana " + _semana_pre)
SetVar("v_ruta_registro_coordinacion", _pre_os.path.join(_re, "COORDINACION DE EMBARQUES",
                                                          "Registro de COAS e informes completos " + _anio_pre + ".xlsx"))

# Leer credenciales SMTP desde Config_Rutas.xlsx (hoja CREDENCIALES)
_config_path = _pre_os.path.join(_br, "Rutas", "Config_Rutas.xlsx")
if _pre_os.path.isfile(_config_path):
    try:
        from openpyxl import load_workbook as _pre_lw
        _wb_cfg = _pre_lw(_config_path, data_only=True)
        if "CREDENCIALES" in _wb_cfg.sheetnames:
            _ws_cred = _wb_cfg["CREDENCIALES"]
            for _r in _ws_cred.iter_rows(min_row=2, values_only=True):
                if not _r or not _r[0]:
                    continue
                _k = str(_r[0]).strip()
                _v = str(_r[1] or "").strip()
                if _k == "SMTP_HOST" and _v: SetVar("v_smtp_host", _v)
                elif _k == "SMTP_PORT" and _v: SetVar("v_smtp_port", _v)
                elif _k == "SMTP_USER" and _v: SetVar("v_smtp_user", _v)
                elif _k == "SMTP_PASS" and _v: SetVar("v_smtp_pass", _v)
                elif _k == "SMTP_SSL" and _v: SetVar("v_smtp_ssl", _v.lower())
    except Exception:
        pass
# Path para libs locales del bot (pypdf, etc. instalados sin admin)
try:
    import sys as _sys_pre
    _libs_pre = _pre_os.path.join(_br, 'libs')
    if _pre_os.path.isdir(_libs_pre) and _libs_pre not in _sys_pre.path:
        _sys_pre.path.insert(0, _libs_pre)
except Exception:
    pass
# FIN AUTODETECT_PROLOGUE

import os
from datetime import datetime

PROCESO = "2_Embarques"
EMPRESA_OBJETIVO = "RIVERSIDE NATURAL FOODS, LTD."

def _log(msg):
    import os as _os
    from datetime import datetime as _dt
    # Resolver log path inline (Rocketbot no ve funciones helper desde otras funciones)
    rp = GetVar("v_orq_log_path") or GetVar("v_ruta_log_actual") or ""
    if not (rp and rp != "ERROR_NOT_VAR" and _os.path.isabs(rp) and _os.path.isdir(_os.path.dirname(rp))):
        rp = ""
        root_emb = GetVar("v_root_embarques") or GetVar("v_ruta_base") or ""
        if root_emb and root_emb != "ERROR_NOT_VAR":
            logs_dir = _os.path.join(root_emb, "0. Bot", "Logs")
            if _os.path.isdir(logs_dir):
                try:
                    logs = [f for f in _os.listdir(logs_dir) if f.startswith("LOG - ") and f.endswith(".txt")]
                    if logs:
                        rp = _os.path.join(logs_dir, sorted(logs)[-1])
                except Exception:
                    pass
    if rp:
        try:
            with open(rp, "a", encoding="utf-8") as f:
                f.write(_dt.now().strftime("%H:%M:%S") + "|" + "2_Embarques" + "|" + msg + chr(10))
        except Exception:
            pass

def _getv(name, default=""):
    v = GetVar(name)
    if not v or v == "ERROR_NOT_VAR":
        return default
    return v

def _leer_txt(path):
    import os as _os
    items = []
    if path and _os.path.exists(path):
        try:
            with open(path, "r", encoding="utf-8") as f:
                for linea in f:
                    s = linea.strip()
                    if s and not s.startswith("#"):
                        items.append(s)
        except Exception:
            pass
    return items

def _hoja_mes_calc(anio_i, semana):
    from datetime import date as _date, timedelta as _td
    MESES = ["ENE","FEB","MAR","ABR","MAY","JUN","JUL","AGO","SEP","OCT","NOV","DIC"]
    jan4 = _date(anio_i, 1, 4)
    lunes_s1 = jan4 - _td(days=jan4.weekday())
    lunes = lunes_s1 + _td(weeks=semana - 1)
    return MESES[lunes.month - 1]

# ============ FLUJO PRINCIPAL ============

_log("Inicio bot 2_Embarques")

ruta_excel = _getv("v_ruta_excel")
if not ruta_excel:
    _log("SIN_DATA | No hay Excel semanal detectado (v_ruta_excel vacio)")
    SetVar("v_total_filas", "0")
else:
    _log("Excel a procesar: " + ruta_excel)
    anio = _getv("v_anio_semana", str(datetime.now().isocalendar()[0]))
    try:
        anio_i = int(anio)
    except Exception:
        anio_i = datetime.now().isocalendar()[0]
        anio = str(anio_i)
    try:
        semana = int(_getv("v_numero_semana_trabajo", str(datetime.now().isocalendar()[1])))
    except Exception:
        semana = datetime.now().isocalendar()[1]
    hoja_mes = _hoja_mes_calc(anio_i, semana)
    SetVar("v_mes_semana", hoja_mes)
    _log("Semana " + str(semana) + " anio " + str(anio_i) + " -> hoja " + hoja_mes)

    root_embarques = _getv("v_root_embarques") or _getv("v_ruta_base")
    carpeta_coord = _getv("v_carpeta_coordinacion_embarques", "COORDINACION DE EMBARQUES")
    prefijo_reg = _getv("v_nombre_registro_prefijo", "Registro de COAS e informes completos ")
    registro_path = _getv("v_ruta_registro_coordinacion")
    if not registro_path and root_embarques:
        registro_path = os.path.join(root_embarques, carpeta_coord, prefijo_reg + anio + ".xlsx")
        SetVar("v_ruta_registro_coordinacion", registro_path)
    _log("Registro path: " + (registro_path or "VACIO"))

    filas = []
    # Leer TODOS los xlsx de la carpeta semana (no solo v_ruta_excel)
    ruta_semana = _getv("v_ruta_semana")
    if not ruta_semana:
        ruta_semana = os.path.dirname(ruta_excel) if ruta_excel else ""
    excel_list = []
    if ruta_semana and os.path.isdir(ruta_semana):
        try:
            for name in sorted(os.listdir(ruta_semana)):
                if name.lower().endswith(".xlsx") and not name.startswith("~$"):
                    excel_list.append(os.path.join(ruta_semana, name))
        except Exception as e:
            _log("ERROR listdir semana: " + str(e))
    if not excel_list and ruta_excel:
        excel_list = [ruta_excel]
    _log("Excel files a procesar: " + str(len(excel_list)))
    for ex_path in excel_list:
        _log("  -> " + os.path.basename(ex_path))
    try:
        from openpyxl import load_workbook
        total_read_all = 0
        for ex_path in excel_list:
            try:
                wb_src = load_workbook(ex_path, data_only=True)
                ws_src = wb_src.active
                total_read_file = 0
                riverside_file = 0
                for row in ws_src.iter_rows(min_row=2, values_only=True):
                    total_read_file += 1
                    if not row or len(row) < 9:
                        continue
                    cliente = str(row[5] or "").strip()
                    if EMPRESA_OBJETIVO.upper() not in cliente.upper():
                        continue
                    emb = str(row[0] or "").strip()
                    if not emb:
                        continue
                    fec = row[3]
                    ped = str(row[8] or "").strip()
                    sem = str(row[2] or "").strip() if len(row) > 2 else ""
                    fec_str = fec.strftime("%d/%m/%Y") if hasattr(fec, "strftime") else str(fec or "").strip()
                    filas.append((cliente, emb, fec_str, ped, sem))
                    riverside_file += 1
                total_read_all += total_read_file
                _log("  " + os.path.basename(ex_path) + ": " + str(total_read_file) + " filas, " + str(riverside_file) + " RIVERSIDE")
            except Exception as e_file:
                _log("ERROR al leer " + os.path.basename(ex_path) + ": " + str(e_file))
        # Dedup entre xlsx por clave embarque|pedido
        visto = set()
        filas_unicas = []
        for (cli, emb, fec, ped, sem) in filas:
            k = emb + "|" + ped
            if k in visto:
                continue
            visto.add(k)
            filas_unicas.append((cli, emb, fec, ped, sem))
        _log("Total filas leidas cross-files: " + str(total_read_all) + " | RIVERSIDE matches: " + str(len(filas)) + " | unicas: " + str(len(filas_unicas)))
        filas = filas_unicas
    except Exception as e:
        _log("ERROR general al leer xlsx: " + str(e))
        SetVar("v_total_filas", "0")
        filas = []

    SetVar("v_total_filas", str(len(filas)))

    nuevos = []
    duplicados = []
    if filas and registro_path and os.path.exists(registro_path):
        try:
            from openpyxl import load_workbook
            wb_reg = load_workbook(registro_path)
            if hoja_mes in wb_reg.sheetnames:
                ws_reg = wb_reg[hoja_mes]
                claves = set()
                for row in ws_reg.iter_rows(min_row=2, values_only=True):
                    if not row or len(row) < 5:
                        continue
                    b = str(row[2] or "").strip()
                    d = str(row[4] or "").strip()
                    if b:
                        claves.add(b + "|" + d)
                _log("Registro " + hoja_mes + ": " + str(len(claves)) + " embarques ya cargados")
                next_row = ws_reg.max_row + 1
                for (cli, emb, fec, ped, sem) in filas:
                    clave = emb + "|" + ped
                    if clave in claves:
                        duplicados.append(emb)
                        continue
                    ws_reg.cell(row=next_row, column=1, value=sem)
                    ws_reg.cell(row=next_row, column=2, value=cli)
                    ws_reg.cell(row=next_row, column=3, value=emb)
                    ws_reg.cell(row=next_row, column=4, value=fec)
                    ws_reg.cell(row=next_row, column=5, value=ped)
                    claves.add(clave)
                    nuevos.append((cli, emb, fec, ped, sem))
                    next_row += 1
                if nuevos:
                    wb_reg.save(registro_path)
                    _log("Registro guardado con " + str(len(nuevos)) + " filas nuevas")
            else:
                _log("ERROR hoja " + hoja_mes + " no existe en Registro")
        except Exception as e:
            _log("ERROR al escribir Registro: " + str(e))
    elif not filas:
        _log("No hay filas de RIVERSIDE que agregar")
    elif not registro_path or not os.path.exists(registro_path):
        _log("ERROR Registro no existe en " + str(registro_path))

    _log("Resumen: total_riverside=" + str(len(filas)) + " nuevos=" + str(len(nuevos)) + " duplicados=" + str(len(duplicados)))

    if nuevos:
        try:
            import smtplib
            from email.mime.multipart import MIMEMultipart
            from email.mime.text import MIMEText
            correos_dir = _getv("v_ruta_correos_coas")
            para = _leer_txt(os.path.join(correos_dir, "Para.txt")) if correos_dir else []
            cc_list = _leer_txt(os.path.join(correos_dir, "CC.txt")) if correos_dir else []
            co_list = _leer_txt(os.path.join(correos_dir, "CO.txt")) if correos_dir else []
            _log("Correos: Para=" + str(len(para)) + " CC=" + str(len(cc_list)) + " CO=" + str(len(co_list)))
            if para:
                host = (GetVar("v_smtp_host") or "smtp.office365.com") if (GetVar("v_smtp_host") not in ("", "ERROR_NOT_VAR")) else "smtp.office365.com"
                _p = GetVar("v_smtp_port"); port = int(_p) if (_p and _p != "ERROR_NOT_VAR") else 587
                user = GetVar("v_smtp_user") if (GetVar("v_smtp_user") not in ("", "ERROR_NOT_VAR")) else ""
                pwd  = GetVar("v_smtp_pass") if (GetVar("v_smtp_pass") not in ("", "ERROR_NOT_VAR")) else ""
                _s = GetVar("v_smtp_ssl"); ssl_mode = (_s.lower() == "true") if (_s and _s != "ERROR_NOT_VAR") else False
                nombre_sem_msg = _getv("v_nombre_semana", "Semana " + str(semana))
                asunto = "COMEX | Nuevos embarques RIVERSIDE - " + nombre_sem_msg
                lineas_txt = chr(10).join("  - Embarque " + e_ + " | PO " + p_ + " | Fecha " + f_ for (_x, e_, f_, p_, _s) in nuevos)
                cuerpo = ("Estimados," + chr(10) + chr(10)
                         + "Se registraron los siguientes embarques nuevos de RIVERSIDE NATURAL FOODS, LTD." + chr(10)
                         + "correspondientes a la " + nombre_sem_msg + " de " + anio + ":" + chr(10) + chr(10)
                         + lineas_txt + chr(10) + chr(10)
                         + "Por favor, marcar COMPLETO en las columnas COAS e Informes del" + chr(10)
                         + "Registro de Coordinacion una vez que los documentos esten disponibles." + chr(10) + chr(10)
                         + "Saludos," + chr(10) + "Bot COMEX MPF" + chr(10))
                msg_obj = MIMEMultipart()
                msg_obj["From"] = user
                msg_obj["To"] = ", ".join(para)
                if cc_list:
                    msg_obj["Cc"] = ", ".join(cc_list)
                msg_obj["Subject"] = asunto
                msg_obj.attach(MIMEText(cuerpo, "plain", "utf-8"))
                destinos = para + cc_list + co_list
                _log("SMTP conectando a " + host + ":" + str(port) + " user=" + (user or "<vacio>"))
                if ssl_mode:
                    smtp_conn = smtplib.SMTP_SSL(host, port, timeout=30)
                else:
                    smtp_conn = smtplib.SMTP(host, port, timeout=30)
                    smtp_conn.starttls()
                if user and pwd:
                    smtp_conn.login(user, pwd)
                smtp_conn.sendmail(user, destinos, msg_obj.as_string())
                smtp_conn.quit()
                _log("Correo enviado OK a " + str(len(destinos)) + " destinatarios")
            else:
                _log("WARN sin destinatarios en Correos/COAS_e_Informes/Para.txt - NO se envia notificacion")
        except Exception as e:
            _log("ERROR envio correo: " + str(e))
    else:
        _log("Sin embarques nuevos -> no se envia notificacion")

_log("Fin bot 2_Embarques")
