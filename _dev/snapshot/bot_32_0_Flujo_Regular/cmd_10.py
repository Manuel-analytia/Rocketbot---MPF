import os, smtplib
from datetime import datetime
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText

PROCESO = "5_Heartbeat"

def _log(msg):
    rp = GetVar("v_orq_log_path") or GetVar("v_ruta_log_actual") or ""
    if rp and rp != "ERROR_NOT_VAR" and os.path.isabs(rp):
        try:
            with open(rp, "a", encoding="utf-8") as f:
                f.write(datetime.now().strftime("%H:%M:%S") + "|" + PROCESO + "|" + msg + chr(10))
        except Exception:
            pass

def _getv(name, default=""):
    v = GetVar(name)
    if not v or v == "ERROR_NOT_VAR":
        return default
    return v

def _leer_txt(path):
    items = []
    if path and os.path.exists(path):
        try:
            with open(path, "r", encoding="utf-8") as f:
                for ln in f:
                    s = ln.strip()
                    if s and not s.startswith("#"):
                        items.append(s)
        except Exception:
            pass
    return items

_log("Inicio heartbeat")

# Leer SMTP de Config_Rutas (las vars del modulo son placeholders)
root_emb = _getv("v_root_embarques")
if root_emb:
    config_path = os.path.join(root_emb, "0. Bot", "Rutas", "Config_Rutas.xlsx")
    if os.path.isfile(config_path):
        try:
            from openpyxl import load_workbook
            wb_cfg = load_workbook(config_path, data_only=True)
            if "CREDENCIALES" in wb_cfg.sheetnames:
                for r in wb_cfg["CREDENCIALES"].iter_rows(min_row=2, values_only=True):
                    if not r or not r[0]:
                        continue
                    k = str(r[0]).strip()
                    v = str(r[1] or "").strip()
                    if k == "SMTP_HOST" and v: SetVar("v_smtp_host", v)
                    elif k == "SMTP_PORT" and v: SetVar("v_smtp_port", v)
                    elif k == "SMTP_USER" and v: SetVar("v_smtp_user", v)
                    elif k == "SMTP_PASS" and v: SetVar("v_smtp_pass", v)
                    elif k == "SMTP_SSL" and v: SetVar("v_smtp_ssl", v.lower())
        except Exception as e:
            _log("WARN no se pudo leer Config_Rutas: " + str(e))

# Path a Correos/Estatus
correos_estatus = os.path.join(root_emb, "0. Bot", "Correos", "Estatus") if root_emb else ""
para = _leer_txt(os.path.join(correos_estatus, "Para.txt")) if correos_estatus else []
cc = _leer_txt(os.path.join(correos_estatus, "CC.txt")) if correos_estatus else []
co = _leer_txt(os.path.join(correos_estatus, "CO.txt")) if correos_estatus else []
_log("Destinatarios Estatus: Para=" + str(len(para)) + " CC=" + str(len(cc)) + " CO=" + str(len(co)))

if not para:
    _log("WARN sin destinatarios en Para.txt de Estatus, no se envia heartbeat")
else:
    # Info de la corrida
    inicio = _getv("v_orq_inicio", "?")
    fin = _getv("v_orq_fin", "?")
    duracion = _getv("v_orq_duracion", "?")
    estado_general = _getv("v_orq_estado_general", "?")
    estado_pre = _getv("v_precheck_estado", "?")
    estado_h1 = _getv("v_orq_estado_1", "?")
    estado_h2 = _getv("v_orq_estado_2", "?")
    estado_h3 = _getv("v_orq_estado_3", "?")
    estado_h4 = _getv("v_orq_estado_4", "?")

    cuerpo = ("Estimados," + chr(10) + chr(10)
             + "Se ejecuto el bot COMEX MPF." + chr(10) + chr(10)
             + "Inicio:    " + inicio + chr(10)
             + "Fin:       " + fin + chr(10)
             + "Duracion:  " + duracion + chr(10)
             + "Estado:    " + estado_general + chr(10) + chr(10)
             + "Detalle por proceso:" + chr(10)
             + "  - Precheck:                  " + estado_pre + chr(10)
             + "  - 1_Crear_Carpetas:          " + estado_h1 + chr(10)
             + "  - 2_Embarques:               " + estado_h2 + chr(10)
             + "  - 3_Revision_COAS_Informe:   " + estado_h3 + chr(10)
             + "  - 4_Revision_Invoice:        " + estado_h4 + chr(10) + chr(10)
             + "Este es un correo automatico de notificacion. La ausencia de este" + chr(10)
             + "correo en el horario habitual puede indicar caida del servidor o" + chr(10)
             + "una falla tecnica que requiere revision." + chr(10) + chr(10)
             + "Saludos," + chr(10) + "Bot COMEX MPF" + chr(10))

    asunto = "Bot COMEX MPF - Ejecucion " + estado_general + " - " + datetime.now().strftime("%d/%m/%Y")

    # SMTP
    host = _getv("v_smtp_host", "smtp.office365.com")
    port_raw = _getv("v_smtp_port", "587")
    try:
        port = int(port_raw)
    except Exception:
        port = 587
    user = _getv("v_smtp_user")
    pwd = _getv("v_smtp_pass")
    ssl_raw = _getv("v_smtp_ssl", "false")
    ssl_mode = (ssl_raw.lower() == "true")

    if not user or not pwd:
        _log("ERROR SMTP_USER o SMTP_PASS vacios en Config_Rutas, no se envia heartbeat")
    else:
        try:
            msg = MIMEMultipart()
            msg["From"] = user
            msg["To"] = ", ".join(para)
            if cc:
                msg["Cc"] = ", ".join(cc)
            msg["Subject"] = asunto
            msg.attach(MIMEText(cuerpo, "plain", "utf-8"))
            destinos = para + cc + co
            if ssl_mode:
                smtp = smtplib.SMTP_SSL(host, port, timeout=300)
            else:
                smtp = smtplib.SMTP(host, port, timeout=300)
                smtp.starttls()
            smtp.login(user, pwd)
            smtp.sendmail(user, destinos, msg.as_string())
            smtp.quit()
            _log("Heartbeat enviado a " + str(len(destinos)) + " destinatarios")
        except Exception as e:
            _log("ERROR heartbeat: " + type(e).__name__ + ": " + str(e))

_log("Fin heartbeat")
